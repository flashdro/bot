import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO

HEADERS = ["Номер", "Тип", "Наименование", "Модель", "Серийный №", "Инвентарный №",
           "Организация-собственник", "МОЛ / ЛКСИ", "Статус по БУ", "Статус по ОУ",
           "ФИО", "Email", "Отдел", "Место установки", "Площадка",
           "Дата выдачи", "Примечание", "Проверено"]

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

def export_to_xlsx(rows, highlight=True) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "ИТ-актив"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r.get("asset_number"), r.get("type"), r.get("manufacturer"), r.get("model"),
            r.get("serial"), r.get("inventory_no"), r.get("organization"), r.get("mol"),
            r.get("status_bu") or r.get("status"), r.get("status_ou"),
            r.get("full_name"), r.get("email"), r.get("department"),
            r.get("workplace") or r.get("location"), r.get("equipment_site") or r.get("site"),
            r.get("issued_at"), r.get("note"), r.get("verified_at") or "",
        ])
        if highlight and r["verified_at"]:
            for cell in ws[ws.max_row]:
                cell.fill = GREEN

    for col, w in zip("ABCDEFGHIJKLMNOPQR",
                      [16, 16, 20, 20, 20, 20, 25, 25, 18, 18, 25, 25, 18, 25, 25, 12, 25, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


HEADER_MAP = {
    "номер": "asset_number",
    "тип": "type",
    "производитель": "manufacturer",
    "наименование": "manufacturer",  # значение из нового файла используется как наименование
    "модель": "model",
    "серийный №": "serial", "серийный номер": "serial", "серийный": "serial", "s/n": "serial",
    "инвентарный №": "inventory_no", "инвентарный номер": "inventory_no", "инвентарный": "inventory_no",
    "организация-собственник": "organization", "организация собственник": "organization",
    "организация": "organization",
    "пользователь": "full_name", "фио": "full_name",
    "площадка": "site",
    "место установки": "workplace", "рабочее место": "workplace", "кабинет": "workplace",
    "мол / лкси": "mol", "мол/лкси": "mol", "мол": "mol",
    "статус по бу": "status_bu", "статус по оу": "status_ou", "статус": "status",
    "отдел": "department",
    "дата выдачи": "issued_at",
    "примечание": "note",
    "email": "email",
    "проверено": None,
}

def _parse_user(field):
    if not field:
        return None, None
    field = str(field).strip()
    if not field:
        return None, None
    m = re.match(r'^(.+?)\s*\((.+?)\)\s*$', field)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return field, None

def _clean(value):
    if value is None:
        return None
    v = str(value).strip()
    if v in ("", "-", "—", "б/н", "б\\н", "нет"):
        return None
    return v

def parse_xlsx(file_bytes: bytes):
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    mapping = [HEADER_MAP.get(h) for h in header]
    result = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        item = {}
        for key, val in zip(mapping, r):
            if key:
                item[key] = val
        if "full_name" in item or "email" in item:
            raw = item.get("full_name")
            name, email = _parse_user(raw)
            item["full_name"] = name
            if not item.get("email"):
                item["email"] = email
        for k in list(item.keys()):
            item[k] = _clean(item[k]) if k != "full_name" else item[k]
        if item.get("status_bu") and not item.get("status"):
            item["status"] = item["status_bu"]
        result.append(item)
    return result
